from .file_asset_substrate_support import *


class FileAssetOoxmlExtractionTest(unittest.TestCase):
    def test_extract_docx_text_blocks_preserves_paragraph_blocks(self) -> None:
        result = extract_docx_text_blocks(
            _build_minimal_docx_bytes(paragraphs=("第一段", "第二段")),
            filename="通知附件.docx",
            source_object_key="campus_notices/source/authenticated/attachments/notice-001.docx",
        )

        self.assertEqual(result.text_blocks, ("第一段", "第二段"))
        self.assertEqual(result.provenance.engine_name, "python-docx")
        self.assertEqual(result.descriptor.filename, "通知附件.docx")
        self.assertEqual(
            result.descriptor.source_object_key,
            "campus_notices/source/authenticated/attachments/notice-001.docx",
        )

    def test_extract_xlsx_text_blocks_preserves_shared_strings_and_formula_cache(self) -> None:
        result = extract_xlsx_text_blocks(
            _build_formula_xlsx_bytes(),
            filename="报名汇总.xlsx",
            source_object_key="campus_notices/source/authenticated/attachments/notice-002.xlsx",
        )

        self.assertEqual(result.provenance.engine_name, "spreadsheetml-zip-xml")
        self.assertEqual(result.descriptor.filename, "报名汇总.xlsx")
        self.assertEqual(
            result.descriptor.source_object_key,
            "campus_notices/source/authenticated/attachments/notice-002.xlsx",
        )
        self.assertIn("项目 | 数量 | 总计", result.text_blocks[0])
        self.assertIn("报名人数 | 2 | 2 [formula=SUM(B2:B2)]", result.text_blocks[0])

    def test_build_ooxml_layout_elements_bytes_normalizes_docx_structure(self) -> None:
        from docx import Document

        buffer = BytesIO()
        document = Document()
        document.add_heading("中央民族大学附件标题", level=1)
        document.add_paragraph("第一段正文")
        document.add_paragraph("第二段正文")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = "值"
        table.cell(1, 0).text = "校区"
        table.cell(1, 1).text = "丰台校区"
        document.save(buffer)

        rows = [
            json.loads(line)
            for line in build_ooxml_layout_elements_bytes(
                buffer.getvalue(),
                filename="attachment.docx",
                source_object_key="source/attachment.docx",
            )
            .decode("utf-8")
            .splitlines()
            if line.strip()
        ]

        self.assertEqual(rows[0]["row_type"], "heading")
        self.assertEqual(rows[0]["heading_level"], 1)
        self.assertEqual(rows[0]["text_normalized"], "中央民族大学附件标题")
        self.assertEqual(rows[1]["row_type"], "paragraph")
        self.assertEqual(rows[1]["text_normalized"], "第一段正文")
        self.assertEqual(rows[-1]["row_type"], "table")
        self.assertEqual(rows[-1]["table_row_count"], 2)
        self.assertEqual(rows[-1]["table_column_count"], 2)
        self.assertEqual(rows[-1]["source_object_key"], "source/attachment.docx")

    def test_build_ooxml_layout_elements_bytes_detects_heading_like_normal_paragraphs(self) -> None:
        from docx import Document

        buffer = BytesIO()
        document = Document()
        document.add_paragraph("中央民族大学2021年高校新任辅导员专题网络培训教学计划")
        document.add_paragraph(
            "为深入学习贯彻习近平新时代中国特色社会主义思想，现将教学计划制定如下："
        )
        document.add_paragraph("学习目标")
        document.add_paragraph("帮助新任辅导员提高政治理论水平和岗位适应能力。")
        document.add_paragraph("二、学习时间")
        document.add_paragraph("本次培训学习时间为2021年11月22日至2022年2月22日。")
        document.save(buffer)

        rows = [
            json.loads(line)
            for line in build_ooxml_layout_elements_bytes(
                buffer.getvalue(),
                filename="plan.docx",
                source_object_key="source/plan.docx",
            )
            .decode("utf-8")
            .splitlines()
            if line.strip()
        ]

        self.assertEqual(rows[0]["row_type"], "heading")
        self.assertEqual(rows[0]["heading_level"], 1)
        self.assertEqual(rows[2]["row_type"], "heading")
        self.assertEqual(rows[2]["heading_level"], 2)
        self.assertEqual(rows[4]["row_type"], "heading")
        self.assertEqual(rows[4]["heading_level"], 2)

    def test_build_ooxml_layout_elements_bytes_normalizes_xlsx_with_forward_fill(self) -> None:
        rows = [
            json.loads(line)
            for line in build_ooxml_layout_elements_bytes(
                _build_merged_heading_xlsx_bytes(),
                filename="attachment.xlsx",
                source_object_key="source/attachment.xlsx",
            )
            .decode("utf-8")
            .splitlines()
            if line.strip()
        ]

        self.assertEqual(rows[0]["row_type"], "heading")
        self.assertEqual(rows[0]["text_normalized"], "本科")
        self.assertEqual(rows[1]["row_type"], "heading")
        self.assertEqual(rows[1]["text_normalized"], "中央民族大学2019-2020学年第2学期")
        self.assertEqual(rows[2]["row_type"], "table")
        self.assertEqual(rows[2]["sheet_name"], "本科")
        self.assertEqual(rows[2]["table_row_count"], 3)
        self.assertEqual(rows[2]["table_column_count"], 3)
        first_cell = next(
            cell for cell in rows[2]["table_cells"] if cell["row"] == 3 and cell["col"] == 1
        )
        self.assertEqual(first_cell["text"], "大学英语（四）")
        self.assertTrue(first_cell["is_forward_filled"])
        formula_cell = next(
            cell for cell in rows[2]["table_cells"] if cell["row"] == 3 and cell["col"] == 3
        )
        self.assertEqual(formula_cell["formula"], "SUM(B2:B2)")

    def test_build_ooxml_layout_elements_bytes_only_forward_fills_leading_missing_cells(
        self,
    ) -> None:
        rows = [
            json.loads(line)
            for line in build_ooxml_layout_elements_bytes(
                _build_partial_continuation_xlsx_bytes(),
                filename="continuation.xlsx",
                source_object_key="source/continuation.xlsx",
            )
            .decode("utf-8")
            .splitlines()
            if line.strip()
        ]

        table_row = next(row for row in rows if row["row_type"] == "table")
        row4_col1 = next(
            cell for cell in table_row["table_cells"] if cell["row"] == 3 and cell["col"] == 1
        )
        self.assertEqual(row4_col1["text"], "大学英语（四）")
        self.assertTrue(row4_col1["is_forward_filled"])
        self.assertFalse(
            any(cell["row"] == 3 and cell["col"] == 3 for cell in table_row["table_cells"])
        )

    def test_serialize_file_asset_outcome_emits_repo_owned_extraction_shape(self) -> None:
        result = extract_docx_text_blocks(
            _build_minimal_docx_bytes(paragraphs=("第一段",)),
            filename="通知附件.docx",
            source_object_key="campus_notices/source/authenticated/attachments/notice-003.docx",
        )

        payload = serialize_file_asset_outcome(result)

        self.assertEqual(payload["outcome_type"], "extraction")
        self.assertEqual(payload["descriptor"]["asset_family"], "modern_ooxml")
        self.assertEqual(payload["descriptor"]["filename"], "通知附件.docx")
        self.assertEqual(
            payload["descriptor"]["source_object_key"],
            "campus_notices/source/authenticated/attachments/notice-003.docx",
        )
        self.assertEqual(payload["provenance"]["engine_name"], "python-docx")
        self.assertEqual(payload["provenance"]["warnings"], [])
        self.assertEqual(payload["text_blocks"], ["第一段"])
        self.assertEqual(payload["artifact_refs"], [])
        parsed = parse_file_asset_outcome(payload)
        self.assertEqual(parsed.descriptor, result.descriptor)
        self.assertEqual(parsed.provenance, result.provenance)
        self.assertEqual(parsed.text_blocks, result.text_blocks)
        self.assertEqual(parsed.extraction_strategy, result.extraction_strategy)
        self.assertEqual(parsed.artifact_refs, result.artifact_refs)

    def test_serialize_file_asset_extraction_record_emits_shared_payload_shape(self) -> None:
        outcome = extract_docx_text_blocks(
            _build_minimal_docx_bytes(paragraphs=("第一段",)),
            filename="通知附件.docx",
            source_object_key="campus_notices/source/authenticated/attachments/notice-004.docx",
        )

        payload = serialize_file_asset_extraction_record(
            FileAssetExtractionRecord(
                outcome=outcome,
                metadata={
                    "attachment_ref_id": "attachment-0001",
                    "extraction_class": "text_extractable",
                    "extraction_policy_version": "policy-v1",
                },
            )
        )

        self.assertEqual(payload["schema_version"], "file_asset_extraction_record_v1")
        self.assertEqual(payload["metadata"]["attachment_ref_id"], "attachment-0001")
        self.assertEqual(payload["metadata"]["extraction_class"], "text_extractable")
        self.assertEqual(payload["metadata"]["extraction_policy_version"], "policy-v1")
        self.assertEqual(payload["outcome"]["outcome_type"], "extraction")
        parsed = parse_file_asset_extraction_record(payload)
        self.assertEqual(parsed.metadata, payload["metadata"])
        self.assertEqual(parsed.outcome.descriptor, outcome.descriptor)
        self.assertEqual(parsed.outcome.provenance, outcome.provenance)
        self.assertEqual(parsed.outcome.text_blocks, outcome.text_blocks)
        self.assertEqual(parsed.outcome.extraction_strategy, outcome.extraction_strategy)
        self.assertEqual(parsed.outcome.artifact_refs, outcome.artifact_refs)
