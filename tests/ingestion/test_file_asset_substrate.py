from __future__ import annotations

import hashlib
import json
import os
import tomllib
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from campus_ingestion_core.file_asset_contract import (
    ArtifactProvenance,
    ExtractionOutcome,
    FileAssetArtifactRef,
    FileAssetDescriptor,
    FileAssetExtractionRecord,
    NormalizationOutcome,
    parse_file_asset_extraction_record,
    parse_file_asset_outcome,
    serialize_file_asset_extraction_record,
    serialize_file_asset_outcome,
)
from campus_ingestion_core.file_asset_legacy_office import (
    extract_legacy_office_text_blocks,
    normalize_legacy_office_asset,
)
from campus_ingestion_core.file_asset_ooxml import (
    build_ooxml_layout_elements_bytes,
    extract_docx_text_blocks,
    extract_xlsx_text_blocks,
)
from campus_ingestion_core.file_asset_routing import (
    RawAssetTypedMetadata,
    infer_raw_asset_typed_metadata,
    refine_raw_asset_typed_metadata_with_ocr_evidence,
    select_adjudicated_image_routing,
    select_online_ocr_routing,
)
from campus_ingestion_core.file_asset_text_extractable import (
    extract_text_extractable_file_asset,
)


def _build_minimal_docx_bytes(*, paragraphs: tuple[str, ...]) -> bytes:
    document_xml = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">',
        "<w:body>",
    ]
    for paragraph in paragraphs:
        document_xml.append(f"<w:p><w:r><w:t>{paragraph}</w:t></w:r></w:p>")
    document_xml.extend(["<w:sectPr/>", "</w:body>", "</w:document>"])

    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>
"""
    root_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>
"""

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("word/document.xml", "\n".join(document_xml))
    return buffer.getvalue()


def _build_formula_xlsx_bytes() -> bytes:
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
</Types>
"""
    root_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>
"""
    workbook = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets>
    <sheet name="预算表" sheetId="1" r:id="rId1"/>
  </sheets>
</workbook>
"""
    workbook_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
</Relationships>
"""
    shared_strings = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="3" uniqueCount="3">
  <si><t>项目</t></si>
  <si><t>数量</t></si>
  <si><t>总计</t></si>
</sst>
"""
    worksheet = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetData>
    <row r="1">
      <c r="A1" t="s"><v>0</v></c>
      <c r="B1" t="s"><v>1</v></c>
      <c r="C1" t="s"><v>2</v></c>
    </row>
    <row r="2">
      <c r="A2" t="inlineStr"><is><t>报名人数</t></is></c>
      <c r="B2"><v>2</v></c>
      <c r="C2"><f>SUM(B2:B2)</f><v>2</v></c>
    </row>
  </sheetData>
</worksheet>
"""
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        archive.writestr("xl/sharedStrings.xml", shared_strings)
        archive.writestr("xl/worksheets/sheet1.xml", worksheet)
    return buffer.getvalue()


def _build_merged_heading_xlsx_bytes() -> bytes:
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
</Types>
"""
    root_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>
"""
    workbook = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets>
    <sheet name="本科" sheetId="1" r:id="rId1"/>
  </sheets>
</workbook>
"""
    workbook_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
</Relationships>
"""
    shared_strings = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="8" uniqueCount="8">
  <si><t>本科</t></si>
  <si><t>中央民族大学2019-2020学年第2学期</t></si>
  <si><t>课程名称</t></si>
  <si><t>专业班级</t></si>
  <si><t>总计</t></si>
  <si><t>大学英语（四）</t></si>
  <si><t>19A1</t></si>
  <si><t>19B1</t></si>
</sst>
"""
    worksheet = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetData>
    <row r="1">
      <c r="A1" t="s"><v>0</v></c>
    </row>
    <row r="2">
      <c r="A2" t="s"><v>1</v></c>
    </row>
    <row r="3">
      <c r="A3" t="s"><v>2</v></c>
      <c r="B3" t="s"><v>3</v></c>
      <c r="C3" t="s"><v>4</v></c>
    </row>
    <row r="4">
      <c r="A4" t="s"><v>5</v></c>
      <c r="B4" t="s"><v>6</v></c>
      <c r="C4"><f>SUM(B2:B2)</f><v>2</v></c>
    </row>
    <row r="5">
      <c r="B5" t="s"><v>7</v></c>
      <c r="C5"><v>3</v></c>
    </row>
  </sheetData>
  <mergeCells count="2">
    <mergeCell ref="A1:C1"/>
    <mergeCell ref="A2:C2"/>
  </mergeCells>
</worksheet>
"""
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        archive.writestr("xl/sharedStrings.xml", shared_strings)
        archive.writestr("xl/worksheets/sheet1.xml", worksheet)
    return buffer.getvalue()


def _build_partial_continuation_xlsx_bytes() -> bytes:
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
</Types>
"""
    root_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>
"""
    workbook = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets>
    <sheet name="本科" sheetId="1" r:id="rId1"/>
  </sheets>
</workbook>
"""
    workbook_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
</Relationships>
"""
    shared_strings = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="10" uniqueCount="10">
  <si><t>考试安排</t></si>
  <si><t>课程名称</t></si>
  <si><t>专业班级</t></si>
  <si><t>主监考</t></si>
  <si><t>备注</t></si>
  <si><t>大学英语（四）</t></si>
  <si><t>19A1</t></si>
  <si><t>鲁雯</t></si>
  <si><t>第一行备注</t></si>
  <si><t>19B1</t></si>
</sst>
"""
    worksheet = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetData>
    <row r="1">
      <c r="A1" t="s"><v>0</v></c>
    </row>
    <row r="2">
      <c r="A2" t="s"><v>1</v></c>
      <c r="B2" t="s"><v>2</v></c>
      <c r="C2" t="s"><v>3</v></c>
      <c r="D2" t="s"><v>4</v></c>
    </row>
    <row r="3">
      <c r="A3" t="s"><v>5</v></c>
      <c r="B3" t="s"><v>6</v></c>
      <c r="C3" t="s"><v>7</v></c>
      <c r="D3" t="s"><v>8</v></c>
    </row>
    <row r="4">
      <c r="B4" t="s"><v>9</v></c>
      <c r="D4" t="inlineStr"><is><t>第二行备注</t></is></c>
    </row>
  </sheetData>
</worksheet>
"""
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        archive.writestr("xl/sharedStrings.xml", shared_strings)
        archive.writestr("xl/worksheets/sheet1.xml", worksheet)
    return buffer.getvalue()


class FileAssetOoxmlExtractionTest(unittest.TestCase):
    def test_extract_docx_text_blocks_preserves_paragraph_blocks(self) -> None:
        result = extract_docx_text_blocks(
            _build_minimal_docx_bytes(paragraphs=("第一段", "第二段")),
            filename="通知附件.docx",
            source_object_key="campus_notices/source/authenticated/attachments/notice-001.docx",
        )

        self.assertEqual(result.text_blocks, ("第一段", "第二段"))
        self.assertEqual(result.provenance.engine_name, "python-docx")
        self.assertEqual(result.descriptor.filename, "通知附件.docx")
        self.assertEqual(
            result.descriptor.source_object_key,
            "campus_notices/source/authenticated/attachments/notice-001.docx",
        )

    def test_extract_xlsx_text_blocks_preserves_shared_strings_and_formula_cache(self) -> None:
        result = extract_xlsx_text_blocks(
            _build_formula_xlsx_bytes(),
            filename="报名汇总.xlsx",
            source_object_key="campus_notices/source/authenticated/attachments/notice-002.xlsx",
        )

        self.assertEqual(result.provenance.engine_name, "spreadsheetml-zip-xml")
        self.assertEqual(result.descriptor.filename, "报名汇总.xlsx")
        self.assertEqual(
            result.descriptor.source_object_key,
            "campus_notices/source/authenticated/attachments/notice-002.xlsx",
        )
        self.assertIn("项目 | 数量 | 总计", result.text_blocks[0])
        self.assertIn("报名人数 | 2 | 2 [formula=SUM(B2:B2)]", result.text_blocks[0])

    def test_build_ooxml_layout_elements_bytes_normalizes_docx_structure(self) -> None:
        from docx import Document

        buffer = BytesIO()
        document = Document()
        document.add_heading("中央民族大学附件标题", level=1)
        document.add_paragraph("第一段正文")
        document.add_paragraph("第二段正文")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = "值"
        table.cell(1, 0).text = "校区"
        table.cell(1, 1).text = "丰台校区"
        document.save(buffer)

        rows = [
            json.loads(line)
            for line in build_ooxml_layout_elements_bytes(
                buffer.getvalue(),
                filename="attachment.docx",
                source_object_key="source/attachment.docx",
            )
            .decode("utf-8")
            .splitlines()
            if line.strip()
        ]

        self.assertEqual(rows[0]["row_type"], "heading")
        self.assertEqual(rows[0]["heading_level"], 1)
        self.assertEqual(rows[0]["text_normalized"], "中央民族大学附件标题")
        self.assertEqual(rows[1]["row_type"], "paragraph")
        self.assertEqual(rows[1]["text_normalized"], "第一段正文")
        self.assertEqual(rows[-1]["row_type"], "table")
        self.assertEqual(rows[-1]["table_row_count"], 2)
        self.assertEqual(rows[-1]["table_column_count"], 2)
        self.assertEqual(rows[-1]["source_object_key"], "source/attachment.docx")

    def test_build_ooxml_layout_elements_bytes_detects_heading_like_normal_paragraphs(self) -> None:
        from docx import Document

        buffer = BytesIO()
        document = Document()
        document.add_paragraph("中央民族大学2021年高校新任辅导员专题网络培训教学计划")
        document.add_paragraph(
            "为深入学习贯彻习近平新时代中国特色社会主义思想，现将教学计划制定如下："
        )
        document.add_paragraph("学习目标")
        document.add_paragraph("帮助新任辅导员提高政治理论水平和岗位适应能力。")
        document.add_paragraph("二、学习时间")
        document.add_paragraph("本次培训学习时间为2021年11月22日至2022年2月22日。")
        document.save(buffer)

        rows = [
            json.loads(line)
            for line in build_ooxml_layout_elements_bytes(
                buffer.getvalue(),
                filename="plan.docx",
                source_object_key="source/plan.docx",
            )
            .decode("utf-8")
            .splitlines()
            if line.strip()
        ]

        self.assertEqual(rows[0]["row_type"], "heading")
        self.assertEqual(rows[0]["heading_level"], 1)
        self.assertEqual(rows[2]["row_type"], "heading")
        self.assertEqual(rows[2]["heading_level"], 2)
        self.assertEqual(rows[4]["row_type"], "heading")
        self.assertEqual(rows[4]["heading_level"], 2)

    def test_build_ooxml_layout_elements_bytes_normalizes_xlsx_with_forward_fill(self) -> None:
        rows = [
            json.loads(line)
            for line in build_ooxml_layout_elements_bytes(
                _build_merged_heading_xlsx_bytes(),
                filename="attachment.xlsx",
                source_object_key="source/attachment.xlsx",
            )
            .decode("utf-8")
            .splitlines()
            if line.strip()
        ]

        self.assertEqual(rows[0]["row_type"], "heading")
        self.assertEqual(rows[0]["text_normalized"], "本科")
        self.assertEqual(rows[1]["row_type"], "heading")
        self.assertEqual(rows[1]["text_normalized"], "中央民族大学2019-2020学年第2学期")
        self.assertEqual(rows[2]["row_type"], "table")
        self.assertEqual(rows[2]["sheet_name"], "本科")
        self.assertEqual(rows[2]["table_row_count"], 3)
        self.assertEqual(rows[2]["table_column_count"], 3)
        first_cell = next(
            cell for cell in rows[2]["table_cells"] if cell["row"] == 3 and cell["col"] == 1
        )
        self.assertEqual(first_cell["text"], "大学英语（四）")
        self.assertTrue(first_cell["is_forward_filled"])
        formula_cell = next(
            cell for cell in rows[2]["table_cells"] if cell["row"] == 3 and cell["col"] == 3
        )
        self.assertEqual(formula_cell["formula"], "SUM(B2:B2)")

    def test_build_ooxml_layout_elements_bytes_only_forward_fills_leading_missing_cells(
        self,
    ) -> None:
        rows = [
            json.loads(line)
            for line in build_ooxml_layout_elements_bytes(
                _build_partial_continuation_xlsx_bytes(),
                filename="continuation.xlsx",
                source_object_key="source/continuation.xlsx",
            )
            .decode("utf-8")
            .splitlines()
            if line.strip()
        ]

        table_row = next(row for row in rows if row["row_type"] == "table")
        row4_col1 = next(
            cell for cell in table_row["table_cells"] if cell["row"] == 3 and cell["col"] == 1
        )
        self.assertEqual(row4_col1["text"], "大学英语（四）")
        self.assertTrue(row4_col1["is_forward_filled"])
        self.assertFalse(
            any(cell["row"] == 3 and cell["col"] == 3 for cell in table_row["table_cells"])
        )

    def test_serialize_file_asset_outcome_emits_repo_owned_extraction_shape(self) -> None:
        result = extract_docx_text_blocks(
            _build_minimal_docx_bytes(paragraphs=("第一段",)),
            filename="通知附件.docx",
            source_object_key="campus_notices/source/authenticated/attachments/notice-003.docx",
        )

        payload = serialize_file_asset_outcome(result)

        self.assertEqual(payload["outcome_type"], "extraction")
        self.assertEqual(payload["descriptor"]["asset_family"], "modern_ooxml")
        self.assertEqual(payload["descriptor"]["filename"], "通知附件.docx")
        self.assertEqual(
            payload["descriptor"]["source_object_key"],
            "campus_notices/source/authenticated/attachments/notice-003.docx",
        )
        self.assertEqual(payload["provenance"]["engine_name"], "python-docx")
        self.assertEqual(payload["provenance"]["warnings"], [])
        self.assertEqual(payload["text_blocks"], ["第一段"])
        self.assertEqual(payload["artifact_refs"], [])
        parsed = parse_file_asset_outcome(payload)
        self.assertEqual(parsed.descriptor, result.descriptor)
        self.assertEqual(parsed.provenance, result.provenance)
        self.assertEqual(parsed.text_blocks, result.text_blocks)
        self.assertEqual(parsed.extraction_strategy, result.extraction_strategy)
        self.assertEqual(parsed.artifact_refs, result.artifact_refs)

    def test_serialize_file_asset_extraction_record_emits_shared_payload_shape(self) -> None:
        outcome = extract_docx_text_blocks(
            _build_minimal_docx_bytes(paragraphs=("第一段",)),
            filename="通知附件.docx",
            source_object_key="campus_notices/source/authenticated/attachments/notice-004.docx",
        )

        payload = serialize_file_asset_extraction_record(
            FileAssetExtractionRecord(
                outcome=outcome,
                metadata={
                    "attachment_ref_id": "attachment-0001",
                    "extraction_class": "text_extractable",
                    "extraction_policy_version": "policy-v1",
                },
            )
        )

        self.assertEqual(payload["schema_version"], "file_asset_extraction_record_v1")
        self.assertEqual(payload["metadata"]["attachment_ref_id"], "attachment-0001")
        self.assertEqual(payload["metadata"]["extraction_class"], "text_extractable")
        self.assertEqual(payload["metadata"]["extraction_policy_version"], "policy-v1")
        self.assertEqual(payload["outcome"]["outcome_type"], "extraction")
        parsed = parse_file_asset_extraction_record(payload)
        self.assertEqual(parsed.metadata, payload["metadata"])
        self.assertEqual(parsed.outcome.descriptor, outcome.descriptor)
        self.assertEqual(parsed.outcome.provenance, outcome.provenance)
        self.assertEqual(parsed.outcome.text_blocks, outcome.text_blocks)
        self.assertEqual(parsed.outcome.extraction_strategy, outcome.extraction_strategy)
        self.assertEqual(parsed.outcome.artifact_refs, outcome.artifact_refs)


class FileAssetTextExtractableDispatcherTest(unittest.TestCase):
    def test_extract_text_extractable_file_asset_routes_doc_to_shared_legacy_seam(self) -> None:
        expected = ExtractionOutcome(
            descriptor=FileAssetDescriptor(
                asset_family="legacy_binary_office",
                filename="制度正文.doc",
                media_type="application/msword",
                source_object_key="source/notice-001.doc",
                source_sha256="hash",
            ),
            provenance=ArtifactProvenance(
                engine_name="libreoffice-soffice+wordprocessingml-zip-xml",
                engine_version="LibreOffice 25.2.1.2;wordprocessingml-zip-xml-v1",
            ),
            text_blocks=("制度正文第一段",),
            extraction_strategy="legacy_office_conversion_then_docx_text_blocks",
        )

        with patch(
            "campus_ingestion_core.file_asset_text_extractable.extract_legacy_office_text_blocks",
            return_value=expected,
        ) as extract_legacy_office:
            result = extract_text_extractable_file_asset(
                b"legacy-doc",
                filename="制度正文.doc",
                media_type="application/msword",
                source_object_key="source/notice-001.doc",
            )

        self.assertIs(result, expected)
        extract_legacy_office.assert_called_once_with(
            b"legacy-doc",
            filename="制度正文.doc",
            media_type="application/msword",
            source_object_key="source/notice-001.doc",
        )

    def test_extract_text_extractable_file_asset_routes_xls_to_shared_legacy_seam(self) -> None:
        expected = ExtractionOutcome(
            descriptor=FileAssetDescriptor(
                asset_family="legacy_binary_office",
                filename="报名表.xls",
                media_type="application/vnd.ms-excel",
                source_object_key="source/notice-001.xls",
                source_sha256="hash",
            ),
            provenance=ArtifactProvenance(
                engine_name="libreoffice-soffice+spreadsheetml-zip-xml",
                engine_version="LibreOffice 25.2.1.2;spreadsheetml-zip-xml-v1",
            ),
            text_blocks=("Sheet1\n姓名 | 张三",),
            extraction_strategy="legacy_office_conversion_then_xlsx_structured_blocks",
        )

        with patch(
            "campus_ingestion_core.file_asset_text_extractable.extract_legacy_office_text_blocks",
            return_value=expected,
        ) as extract_legacy_office:
            result = extract_text_extractable_file_asset(
                b"legacy-xls",
                filename="报名表.xls",
                media_type="application/vnd.ms-excel",
                source_object_key="source/notice-001.xls",
            )

        self.assertIs(result, expected)
        extract_legacy_office.assert_called_once_with(
            b"legacy-xls",
            filename="报名表.xls",
            media_type="application/vnd.ms-excel",
            source_object_key="source/notice-001.xls",
        )


class FileAssetHardeningContractTest(unittest.TestCase):
    def test_ingestion_core_dependencies_do_not_reintroduce_xlrd(self) -> None:
        pyproject_path = (
            Path(__file__).resolve().parents[2] / "packages/ingestion-core/pyproject.toml"
        )
        payload = tomllib.loads(pyproject_path.read_text(encoding="utf-8"))

        dependencies = tuple(str(item) for item in payload["project"]["dependencies"])
        self.assertFalse(any(dep.startswith("xlrd") for dep in dependencies))

    def test_forward_substrate_does_not_keep_retired_compatibility_shims(self) -> None:
        root = Path(__file__).resolve().parents[2]

        self.assertFalse(
            (
                root / "packages/ingestion-core/src/campus_ingestion_core/raw_asset_typing.py"
            ).exists()
        )
        self.assertFalse(
            (
                root / "packages/ingestion-core/src/campus_ingestion_core/online_ocr_routing.py"
            ).exists()
        )


class FileAssetRoutingSubstrateTest(unittest.TestCase):
    def test_infer_raw_asset_typed_metadata_classifies_pdf_as_document_pdf(self) -> None:
        result = infer_raw_asset_typed_metadata(
            source_path=Path("中央民族大学培养方案.pdf"),
            file_bytes=b"%PDF-1.7 sample fixture",
        )

        self.assertEqual(result.asset_kind, "document")
        self.assertEqual(result.asset_subkind, "pdf")
        self.assertEqual(result.file_extension, ".pdf")

    def test_select_online_ocr_routing_prefers_mineru_for_pdf(self) -> None:
        decision = select_online_ocr_routing(
            asset_kind="document",
            asset_subkind="pdf",
            mineru_available=True,
        )

        self.assertTrue(decision.canonical_allowed)
        self.assertEqual(decision.route_id, "pdf_mineru_primary_with_paddle_fallback")
        self.assertEqual(decision.canonical_provider, "mineru_api")
        self.assertEqual(decision.fallback_provider, "paddleocr_async_jobs")

    def test_refined_spatial_image_stays_fail_closed_after_dual_run(self) -> None:
        initial_typed_metadata = RawAssetTypedMetadata(
            content_type="image/png",
            file_extension=".png",
            byte_size=2048,
            page_count=1,
            image_width=1920,
            image_height=1080,
            asset_kind="image",
            asset_subkind="generic_image",
            classification_confidence=1.0,
        )
        candidate_summary = {
            "provider": "mineru_api",
            "usable_output": True,
            "structure_signals": {
                "table_row_count": 0,
                "nonempty_line_count": 2,
                "image_reference_count": 1,
            },
        }
        baseline_summary = {
            "provider": "paddleocr_async_jobs",
            "usable_output": True,
            "structure_signals": {
                "table_row_count": 0,
                "nonempty_line_count": 1,
                "image_reference_count": 0,
            },
        }
        adjudication = {"winner": "candidate", "reason_codes": []}

        refined = refine_raw_asset_typed_metadata_with_ocr_evidence(
            typed_metadata=initial_typed_metadata,
            source_name="campus-map.png",
            candidate_summary=candidate_summary,
            baseline_summary=baseline_summary,
            adjudication=adjudication,
        )
        decision = select_adjudicated_image_routing(
            refined_asset_subkind=refined.asset_subkind,
            candidate_summary=candidate_summary,
            baseline_summary=baseline_summary,
        )

        self.assertEqual(refined.asset_subkind, "spatial_image")
        self.assertFalse(decision.canonical_allowed)
        self.assertEqual(decision.route_id, "spatial_image_deferred")


class LegacyOfficeNormalizationSubstrateTest(unittest.TestCase):
    def test_normalize_legacy_office_asset_converts_xls_to_xlsx_with_provenance(self) -> None:
        with TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            input_path = temp_root / "legacy-score-sheet.xls"
            input_bytes = b"legacy-xls-binary"
            input_path.write_bytes(input_bytes)
            output_dir = temp_root / "normalized"
            output_path = output_dir / "legacy-score-sheet.xlsx"
            normalized_bytes = b"normalized-xlsx-binary"

            result = normalize_legacy_office_asset(
                input_path=input_path,
                output_dir=output_dir,
                source_object_key="attachments/example/legacy-score-sheet.xls",
                source_mime_type="application/vnd.ms-excel",
                normalized_object_key="campus_notices/operations/normalized/legacy-score-sheet.xlsx",
                conversion_client=SimpleNamespace(
                    convert=lambda request: SimpleNamespace(
                        target_format="xlsx",
                        normalized_media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        normalized_sha256=hashlib.sha256(normalized_bytes).hexdigest(),
                        normalized_bytes=normalized_bytes,
                        engine_name="jodconverter+libreoffice",
                        engine_version="7.6.7",
                        warnings=(),
                    )
                ),
            )

        self.assertEqual(result.source_extension, ".xls")
        self.assertIsInstance(result, NormalizationOutcome)
        self.assertEqual(result.normalized_extension, ".xlsx")
        self.assertEqual(result.output_path, str(output_path))
        self.assertEqual(result.descriptor.asset_family, "legacy_binary_office")
        self.assertEqual(result.descriptor.filename, "legacy-score-sheet.xls")
        self.assertEqual(
            result.descriptor.source_object_key,
            "attachments/example/legacy-score-sheet.xls",
        )
        self.assertEqual(
            result.artifact_refs,
            (
                FileAssetArtifactRef(
                    role="normalized_object",
                    object_key="campus_notices/operations/normalized/legacy-score-sheet.xlsx",
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
        )
        self.assertEqual(result.provenance.engine_name, "jodconverter+libreoffice")
        self.assertEqual(result.provenance.engine_version, "7.6.7")
        self.assertEqual(
            result.source_sha256,
            hashlib.sha256(input_bytes).hexdigest(),
        )
        self.assertEqual(
            result.normalized_sha256,
            hashlib.sha256(normalized_bytes).hexdigest(),
        )

    def test_extract_legacy_office_text_blocks_converts_xls_via_shared_normalization(self) -> None:
        normalized_bytes = _build_formula_xlsx_bytes()

        result = extract_legacy_office_text_blocks(
            b"legacy-xls-binary",
            filename="legacy-score-sheet.xls",
            media_type="application/vnd.ms-excel",
            source_object_key="attachments/example/legacy-score-sheet.xls",
            conversion_client=SimpleNamespace(
                convert=lambda request: SimpleNamespace(
                    target_format="xlsx",
                    normalized_media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    normalized_sha256=hashlib.sha256(normalized_bytes).hexdigest(),
                    normalized_bytes=normalized_bytes,
                    engine_name="jodconverter+libreoffice",
                    engine_version="7.6.7",
                    warnings=(),
                )
            ),
        )

        self.assertEqual(result.descriptor.asset_family, "legacy_binary_office")
        self.assertEqual(result.descriptor.filename, "legacy-score-sheet.xls")
        self.assertEqual(
            result.descriptor.source_object_key,
            "attachments/example/legacy-score-sheet.xls",
        )
        self.assertEqual(
            result.provenance.engine_name,
            "jodconverter+libreoffice+spreadsheetml-zip-xml",
        )
        self.assertIn("normalized_via_legacy_office_conversion", result.provenance.warnings)
        self.assertIn("项目 | 数量 | 总计", result.text_blocks[0])
        self.assertIn("报名人数 | 2 | 2 [formula=SUM(B2:B2)]", result.text_blocks[0])

    def test_normalize_legacy_office_asset_fails_close_when_service_is_not_configured(self) -> None:
        with TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            input_path = temp_root / "legacy-score-sheet.xls"
            input_path.write_bytes(b"legacy-xls-binary")
            output_dir = temp_root / "normalized"

            with (
                patch.dict(
                    os.environ,
                    {
                        "LEGACY_OFFICE_CONVERSION_BASE_URL": "",
                        "LEGACY_OFFICE_CONVERSION_INTERNAL_SHARED_SECRET": "",
                    },
                ),
                self.assertRaisesRegex(
                    RuntimeError,
                    "legacy Office conversion service is not configured",
                ),
            ):
                normalize_legacy_office_asset(
                    input_path=input_path,
                    output_dir=output_dir,
                )


if __name__ == "__main__":
    unittest.main()
