from __future__ import annotations

import hashlib
import posixpath
import re
import zipfile
from dataclasses import dataclass
from importlib import metadata
from io import BytesIO
from xml.etree import ElementTree

from campus_ingestion_core.file_asset_contract import (
    ArtifactProvenance,
    ExtractionOutcome,
    FileAssetDescriptor,
)
from campus_ingestion_core.source_truth_common import (
    detect_merged_cells,
    infer_list_level_hint,
    infer_table_column_count,
    infer_table_header_rows,
    infer_table_row_count,
    normalize_text,
    serialize_layout_rows,
)

_DOCX_DOCUMENT_PATH = "word/document.xml"
_WORDPROCESSINGML_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_XLSX_WORKBOOK_PATH = "xl/workbook.xml"
_XLSX_WORKBOOK_RELS_PATH = "xl/_rels/workbook.xml.rels"
_XLSX_SHARED_STRINGS_PATH = "xl/sharedStrings.xml"
_SPREADSHEETML_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_PACKAGE_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_OFFICE_DOC_REL_ID_ATTR = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_OOXML_LAYOUT_SCHEMA_VERSION = "ooxml_layout_elements_v1"
_DOCX_SECTION_HEADING_RE = re.compile(r"^[一二三四五六七八九十]+、")
_DOCX_SUBSECTION_HEADING_RE = re.compile(r"^[（(][一二三四五六七八九十0-9]+[）)]")
_DOCX_TERMINAL_PUNCTUATION = ("。", "！", "？", "；", "：", ":", ".", "!", "?", ";")


@dataclass(frozen=True, slots=True)
class OoxmlTextExtractionResult(ExtractionOutcome):
    pass


@dataclass(frozen=True, slots=True)
class _XlsxCellValue:
    text: str
    formula: str | None = None


@dataclass(frozen=True, slots=True)
class _XlsxSheetRow:
    row_number: int
    cells: dict[int, _XlsxCellValue]


def extract_docx_text_blocks(
    file_bytes: bytes,
    *,
    filename: str = "attachment.docx",
    source_object_key: str | None = None,
    media_type: str = _DOCX_MEDIA_TYPE,
) -> OoxmlTextExtractionResult:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("python-docx is required for docx extraction") from exc

    document = Document(BytesIO(file_bytes))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            blocks.append("\n".join(rows))
    if not blocks:
        raise ValueError("docx attachment produced no extractable text")
    return OoxmlTextExtractionResult(
        descriptor=FileAssetDescriptor(
            asset_family="modern_ooxml",
            filename=filename,
            media_type=media_type,
            source_object_key=source_object_key,
            source_sha256=hashlib.sha256(file_bytes).hexdigest(),
        ),
        provenance=ArtifactProvenance(
            engine_name="python-docx",
            engine_version=_read_package_version("python-docx"),
        ),
        text_blocks=tuple(blocks),
        extraction_strategy="docx_structured_blocks",
    )


def build_ooxml_layout_elements_bytes(
    file_bytes: bytes,
    *,
    filename: str,
    source_object_key: str | None = None,
    media_type: str | None = None,
) -> bytes:
    extension = filename.rsplit(".", 1)[-1].strip().lower() if "." in filename else ""
    if extension == "docx":
        rows = _build_docx_layout_rows(
            file_bytes,
            filename=filename,
            source_object_key=source_object_key,
        )
    elif extension == "xlsx":
        rows = _build_xlsx_layout_rows(
            file_bytes,
            filename=filename,
            source_object_key=source_object_key,
        )
    else:
        raise ValueError(f"unsupported ooxml asset extension: {filename}")
    if not rows:
        raise ValueError(f"{filename} produced no OOXML layout rows")
    return serialize_layout_rows(rows)


def extract_xlsx_text_blocks(
    file_bytes: bytes,
    *,
    filename: str = "attachment.xlsx",
    source_object_key: str | None = None,
    media_type: str = _XLSX_MEDIA_TYPE,
) -> OoxmlTextExtractionResult:
    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
            workbook_xml = archive.read(_XLSX_WORKBOOK_PATH)
            workbook_root = ElementTree.fromstring(workbook_xml)
            workbook_relationship_targets = _load_xlsx_workbook_relationship_targets(archive)
            shared_strings = _load_xlsx_shared_strings(archive)
            blocks: list[str] = []
            for sheet in workbook_root.findall(f".//{_SPREADSHEETML_NS}sheet"):
                rel_id = sheet.attrib.get(_OFFICE_DOC_REL_ID_ATTR)
                if not rel_id:
                    continue
                worksheet_target = workbook_relationship_targets.get(rel_id)
                if worksheet_target is None:
                    continue
                worksheet_path = _resolve_xlsx_zip_member_path(
                    base_path=_XLSX_WORKBOOK_PATH,
                    target=worksheet_target,
                )
                sheet_xml = archive.read(worksheet_path)
                sheet_rows = _load_xlsx_sheet_rows(
                    sheet_xml=sheet_xml,
                    shared_strings=shared_strings,
                )
                if not sheet_rows:
                    continue
                sheet_name = (sheet.attrib.get("name") or "").strip()
                blocks.append(_serialize_spreadsheet_sheet_block(sheet_name, sheet_rows))
    except KeyError as exc:
        raise ValueError("xlsx workbook.xml is required") from exc
    if not blocks:
        raise ValueError("xlsx attachment produced no extractable text")
    return OoxmlTextExtractionResult(
        descriptor=FileAssetDescriptor(
            asset_family="modern_ooxml",
            filename=filename,
            media_type=media_type,
            source_object_key=source_object_key,
            source_sha256=hashlib.sha256(file_bytes).hexdigest(),
        ),
        provenance=ArtifactProvenance(
            engine_name="spreadsheetml-zip-xml",
            engine_version="spreadsheetml-zip-xml-v1",
        ),
        text_blocks=tuple(blocks),
        extraction_strategy="xlsx_structured_blocks",
    )


def _build_docx_layout_rows(
    file_bytes: bytes,
    *,
    filename: str,
    source_object_key: str | None,
) -> list[dict[str, object]]:
    try:
        from docx import Document
        from docx.document import Document as DocxDocument
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("python-docx is required for docx extraction") from exc

    document = Document(BytesIO(file_bytes))
    rows: list[dict[str, object]] = []
    row_index = 0

    def iter_block_items(parent: DocxDocument):
        parent_element = parent.element.body
        for child in parent_element.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    block_items = list(iter_block_items(document))
    for block_index, block in enumerate(block_items):
        row_index += 1
        if block.__class__.__name__ == "Paragraph":
            text_raw = block.text
            text_normalized = normalize_text(text_raw)
            if not text_normalized:
                row_index -= 1
                continue
            style_name = str(getattr(getattr(block, "style", None), "name", "") or "").strip()
            row_type = "paragraph"
            block_role = "paragraph"
            heading_level = None
            list_level_hint = infer_list_level_hint(text_normalized)
            style_lower = style_name.lower()
            if style_lower == "title":
                row_type = "heading"
                block_role = "title"
                heading_level = 1
            elif style_lower.startswith("heading"):
                row_type = "heading"
                block_role = "heading"
                trailing = "".join(character for character in style_name if character.isdigit())
                heading_level = int(trailing) if trailing else 1
            else:
                next_paragraph_text = _find_next_docx_paragraph_text(
                    block_items=block_items,
                    start_index=block_index + 1,
                )
                inferred_heading_level = _infer_docx_heading_level(
                    text_normalized=text_normalized,
                    next_paragraph_text=next_paragraph_text,
                    is_first_paragraph=block_index == 0,
                    list_level_hint=list_level_hint,
                )
                if inferred_heading_level is not None:
                    row_type = "heading"
                    block_role = "title" if inferred_heading_level == 1 else "heading"
                    heading_level = inferred_heading_level
                elif "list" in style_lower or list_level_hint is not None:
                    row_type = "list_item"
                    block_role = "list_item"
            rows.append(
                {
                    "schema_version": _OOXML_LAYOUT_SCHEMA_VERSION,
                    "row_id": f"row-{row_index:04d}",
                    "row_index": row_index,
                    "row_type": row_type,
                    "block_role": block_role,
                    "heading_level": heading_level,
                    "list_level_hint": list_level_hint,
                    "text_raw": text_raw,
                    "text_normalized": text_normalized,
                    "source_object_key": source_object_key,
                    "source_format": "docx",
                    "source_filename": filename,
                }
            )
            continue

        table_cells: list[dict[str, object]] = []
        for table_row_index, table_row in enumerate(block.rows, start=1):
            for table_column_index, cell in enumerate(table_row.cells, start=1):
                text = cell.text.strip()
                if not text:
                    continue
                table_cells.append(
                    {
                        "row": table_row_index,
                        "col": table_column_index,
                        "text": text,
                        "is_header": table_row_index == 1,
                    }
                )
        if not table_cells:
            row_index -= 1
            continue
        rows.append(
            {
                "schema_version": _OOXML_LAYOUT_SCHEMA_VERSION,
                "row_id": f"row-{row_index:04d}",
                "row_index": row_index,
                "row_type": "table",
                "block_role": "table",
                "heading_level": None,
                "text_raw": "",
                "text_normalized": "",
                "table_cells": table_cells,
                "table_header_rows": infer_table_header_rows(table_cells),
                "table_row_count": infer_table_row_count(table_cells),
                "table_column_count": infer_table_column_count(table_cells),
                "table_has_merged_cells": detect_merged_cells(table_cells),
                "source_object_key": source_object_key,
                "source_format": "docx",
                "source_filename": filename,
            }
        )

    return rows


def _build_xlsx_layout_rows(
    file_bytes: bytes,
    *,
    filename: str,
    source_object_key: str | None,
) -> list[dict[str, object]]:
    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
            workbook_xml = archive.read(_XLSX_WORKBOOK_PATH)
            workbook_root = ElementTree.fromstring(workbook_xml)
            workbook_relationship_targets = _load_xlsx_workbook_relationship_targets(archive)
            shared_strings = _load_xlsx_shared_strings(archive)
            rows: list[dict[str, object]] = []
            row_index = 0
            for sheet_position, sheet in enumerate(
                workbook_root.findall(f".//{_SPREADSHEETML_NS}sheet"),
                start=1,
            ):
                rel_id = sheet.attrib.get(_OFFICE_DOC_REL_ID_ATTR)
                if not rel_id:
                    continue
                worksheet_target = workbook_relationship_targets.get(rel_id)
                if worksheet_target is None:
                    continue
                worksheet_path = _resolve_xlsx_zip_member_path(
                    base_path=_XLSX_WORKBOOK_PATH,
                    target=worksheet_target,
                )
                worksheet_root = ElementTree.fromstring(archive.read(worksheet_path))
                sheet_rows = _load_xlsx_sheet_row_payloads(
                    worksheet_root=worksheet_root,
                    shared_strings=shared_strings,
                )
                if not sheet_rows:
                    continue
                heading_rows, table_rows = _partition_xlsx_sheet_rows(sheet_rows)
                sheet_name = normalize_text(str(sheet.attrib.get("name") or ""))
                for heading_offset, sheet_row in enumerate(heading_rows, start=1):
                    heading_text = _extract_single_cell_text(sheet_row.cells)
                    if not heading_text:
                        continue
                    row_index += 1
                    rows.append(
                        {
                            "schema_version": _OOXML_LAYOUT_SCHEMA_VERSION,
                            "row_id": f"sheet-{sheet_position:04d}-row-{row_index:04d}",
                            "row_index": row_index,
                            "row_type": "heading",
                            "block_role": "heading",
                            "heading_level": heading_offset,
                            "text_raw": heading_text,
                            "text_normalized": heading_text,
                            "sheet_name": sheet_name or None,
                            "worksheet_row_number": sheet_row.row_number,
                            "source_object_key": source_object_key,
                            "source_format": "xlsx",
                            "source_filename": filename,
                        }
                    )
                if not table_rows:
                    continue
                row_index += 1
                rows.append(
                    _build_xlsx_table_layout_row(
                        row_index=row_index,
                        sheet_name=sheet_name or None,
                        sheet_rows=table_rows,
                        source_object_key=source_object_key,
                        filename=filename,
                    )
                )
    except KeyError as exc:
        raise ValueError("xlsx workbook.xml is required") from exc
    return rows


def _load_xlsx_sheet_row_payloads(
    *,
    worksheet_root,
    shared_strings: tuple[str, ...],
) -> list[_XlsxSheetRow]:
    rows: list[_XlsxSheetRow] = []
    for row in worksheet_root.findall(f".//{_SPREADSHEETML_NS}row"):
        row_number = int(row.attrib.get("r") or len(rows) + 1)
        cells: dict[int, _XlsxCellValue] = {}
        for cell in row.findall(f"{_SPREADSHEETML_NS}c"):
            value = _load_xlsx_cell_value(cell=cell, shared_strings=shared_strings)
            if value is None:
                continue
            column_index = _xlsx_column_index(cell.attrib.get("r")) + 1
            cells[column_index] = value
        if cells:
            rows.append(_XlsxSheetRow(row_number=row_number, cells=cells))
    return rows


def _partition_xlsx_sheet_rows(
    sheet_rows: list[_XlsxSheetRow],
) -> tuple[list[_XlsxSheetRow], list[_XlsxSheetRow]]:
    heading_rows: list[_XlsxSheetRow] = []
    for index, sheet_row in enumerate(sheet_rows):
        if len(sheet_row.cells) <= 1:
            heading_rows.append(sheet_row)
            continue
        return heading_rows, sheet_rows[index:]
    return heading_rows, []


def _extract_single_cell_text(cells: dict[int, _XlsxCellValue]) -> str | None:
    if len(cells) != 1:
        return None
    value = next(iter(cells.values()))
    text = normalize_text(value.text)
    return text or None


def _build_xlsx_table_layout_row(
    *,
    row_index: int,
    sheet_name: str | None,
    sheet_rows: list[_XlsxSheetRow],
    source_object_key: str | None,
    filename: str,
) -> dict[str, object]:
    max_column = max(
        (max(sheet_row.cells) for sheet_row in sheet_rows if sheet_row.cells), default=0
    )
    table_cells: list[dict[str, object]] = []
    previous_text_by_column: dict[int, str] = {}
    previous_formula_by_column: dict[int, str] = {}
    for table_row_index, sheet_row in enumerate(sheet_rows, start=1):
        explicit_cells = sheet_row.cells
        first_explicit_column = min(explicit_cells) if explicit_cells else None
        continuation_row = (
            table_row_index > 1
            and bool(explicit_cells)
            and first_explicit_column is not None
            and first_explicit_column > 1
        )
        for column_index in range(1, max_column + 1):
            explicit_value = explicit_cells.get(column_index)
            if explicit_value is None:
                if (
                    table_row_index == 1
                    or column_index not in previous_text_by_column
                    or first_explicit_column is None
                    or column_index >= first_explicit_column
                ):
                    continue
                table_cells.append(
                    {
                        "row": table_row_index,
                        "col": column_index,
                        "text": previous_text_by_column[column_index],
                        "formula": previous_formula_by_column.get(column_index),
                        "is_header": table_row_index == 1,
                        "is_forward_filled": True,
                        "worksheet_row_number": sheet_row.row_number,
                    }
                )
                continue
            normalized_text = explicit_value.text.strip()
            formula = explicit_value.formula
            if continuation_row and formula is None and column_index in previous_formula_by_column:
                formula = previous_formula_by_column[column_index]
            cell_payload: dict[str, object] = {
                "row": table_row_index,
                "col": column_index,
                "text": normalized_text,
                "is_header": table_row_index == 1,
                "is_forward_filled": False,
                "worksheet_row_number": sheet_row.row_number,
            }
            if formula is not None:
                cell_payload["formula"] = formula
            table_cells.append(cell_payload)
            if normalized_text:
                previous_text_by_column[column_index] = normalized_text
            if formula is not None:
                previous_formula_by_column[column_index] = formula
    return {
        "schema_version": _OOXML_LAYOUT_SCHEMA_VERSION,
        "row_id": f"row-{row_index:04d}",
        "row_index": row_index,
        "row_type": "table",
        "block_role": "table",
        "heading_level": None,
        "text_raw": "",
        "text_normalized": "",
        "sheet_name": sheet_name,
        "table_cells": table_cells,
        "table_header_rows": infer_table_header_rows(table_cells),
        "table_row_count": infer_table_row_count(table_cells),
        "table_column_count": infer_table_column_count(table_cells),
        "table_has_merged_cells": detect_merged_cells(table_cells),
        "source_object_key": source_object_key,
        "source_format": "xlsx",
        "source_filename": filename,
    }


def _find_next_docx_paragraph_text(
    *,
    block_items: list[object],
    start_index: int,
) -> str | None:
    for block in block_items[start_index:]:
        if block.__class__.__name__ != "Paragraph":
            continue
        text = normalize_text(getattr(block, "text", ""))
        if text:
            return text
    return None


def _infer_docx_heading_level(
    *,
    text_normalized: str,
    next_paragraph_text: str | None,
    is_first_paragraph: bool,
    list_level_hint: int | None,
) -> int | None:
    if not text_normalized:
        return None
    if _DOCX_SECTION_HEADING_RE.match(text_normalized):
        return 2
    if _DOCX_SUBSECTION_HEADING_RE.match(text_normalized):
        return 3
    if list_level_hint is not None:
        return None
    if any(text_normalized.endswith(punctuation) for punctuation in _DOCX_TERMINAL_PUNCTUATION):
        return None
    text_length = len(text_normalized)
    if is_first_paragraph and text_length <= 60:
        return 1
    if next_paragraph_text is not None and text_length <= 12 and len(next_paragraph_text) >= 18:
        return 2
    return None


def _load_xlsx_workbook_relationship_targets(archive: zipfile.ZipFile) -> dict[str, str]:
    try:
        relationships_root = ElementTree.fromstring(archive.read(_XLSX_WORKBOOK_RELS_PATH))
    except KeyError:
        return {}
    return {
        str(node.attrib.get("Id") or "").strip(): str(node.attrib.get("Target") or "").strip()
        for node in relationships_root.findall(f"{_PACKAGE_REL_NS}Relationship")
        if str(node.attrib.get("Id") or "").strip() and str(node.attrib.get("Target") or "").strip()
    }


def _load_xlsx_shared_strings(archive: zipfile.ZipFile) -> tuple[str, ...]:
    try:
        shared_strings_root = ElementTree.fromstring(archive.read(_XLSX_SHARED_STRINGS_PATH))
    except KeyError:
        return ()
    values: list[str] = []
    for node in shared_strings_root.findall(f"{_SPREADSHEETML_NS}si"):
        values.append(
            "".join(
                text_node.text or "" for text_node in node.iter(f"{_SPREADSHEETML_NS}t")
            ).strip()
        )
    return tuple(values)


def _resolve_xlsx_zip_member_path(*, base_path: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_path), target))


def _load_xlsx_sheet_rows(
    *,
    sheet_xml: bytes,
    shared_strings: tuple[str, ...],
) -> tuple[str, ...]:
    root = ElementTree.fromstring(sheet_xml)
    rows: list[str] = []
    for row in root.findall(f".//{_SPREADSHEETML_NS}row"):
        cells: dict[int, _XlsxCellValue] = {}
        for cell in row.findall(f"{_SPREADSHEETML_NS}c"):
            value = _load_xlsx_cell_value(cell=cell, shared_strings=shared_strings)
            if value is None:
                continue
            column_index = _xlsx_column_index(cell.attrib.get("r"))
            cells[column_index] = value
        serialized = _serialize_spreadsheet_row(cells)
        if serialized is not None:
            rows.append(serialized)
    return tuple(rows)


def _load_xlsx_cell_value(
    *,
    cell,
    shared_strings: tuple[str, ...],
) -> _XlsxCellValue | None:
    cell_type = str(cell.attrib.get("t") or "").strip()
    formula_node = cell.find(f"{_SPREADSHEETML_NS}f")
    formula = None
    if formula_node is not None and formula_node.text is not None:
        formula = formula_node.text.strip() or None

    if cell_type == "inlineStr":
        value = "".join(
            text_node.text or "" for text_node in cell.iter(f"{_SPREADSHEETML_NS}t")
        ).strip()
        if not value:
            return None
        return _XlsxCellValue(text=value, formula=formula)

    value_node = cell.find(f"{_SPREADSHEETML_NS}v")
    if value_node is None or value_node.text is None:
        return _XlsxCellValue(text=f"[formula={formula}]", formula=formula) if formula else None
    raw_value = value_node.text.strip()
    if not raw_value:
        return _XlsxCellValue(text=f"[formula={formula}]", formula=formula) if formula else None
    if cell_type == "s":
        try:
            shared_string = shared_strings[int(raw_value)]
        except (IndexError, ValueError):
            text = raw_value
        else:
            text = shared_string or raw_value
        return _XlsxCellValue(text=text, formula=formula)
    if cell_type == "b":
        return _XlsxCellValue(text="TRUE" if raw_value == "1" else "FALSE", formula=formula)
    return _XlsxCellValue(text=raw_value, formula=formula)


def _xlsx_column_index(cell_ref: str | None) -> int:
    if cell_ref is None:
        return 0
    letters = "".join(character for character in cell_ref if character.isalpha()).upper()
    if not letters:
        return 0
    value = 0
    for character in letters:
        value = value * 26 + (ord(character) - ord("A") + 1)
    return value - 1


def _serialize_spreadsheet_sheet_block(sheet_name: str, rows: tuple[str, ...]) -> str:
    if sheet_name:
        return f"{sheet_name}\n" + "\n".join(rows)
    return "\n".join(rows)


def _serialize_spreadsheet_row(cells: dict[int, _XlsxCellValue]) -> str | None:
    if not cells:
        return None
    max_index = max(cells)
    values: list[str] = []
    for index in range(max_index + 1):
        cell = cells.get(index)
        if cell is None:
            values.append("")
            continue
        text = cell.text
        if cell.formula:
            text = f"{text} [formula={cell.formula}]"
        values.append(text)
    return " | ".join(value.strip() for value in values).strip()


def _read_package_version(distribution_name: str) -> str:
    try:
        return metadata.version(distribution_name)
    except metadata.PackageNotFoundError:  # pragma: no cover
        return "unknown"


__all__ = [
    "OoxmlTextExtractionResult",
    "build_ooxml_layout_elements_bytes",
    "extract_docx_text_blocks",
    "extract_xlsx_text_blocks",
]
